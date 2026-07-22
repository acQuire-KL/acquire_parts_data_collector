import argparse, json, re
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from config import Settings
from digikey_client import DigiKeyClient
from manufacturer_resolver import resolve_manufacturer

MFG = {'manufacturer','mfg','mfr','manufacturer name'}
MPN = {'mpn','manufacturer part number','mfg part number','manufacturer_part_number'}

def clean(v): return ' '.join(str(v or '').strip().lower().replace('_',' ').split())
def norm(v): return ''.join(c for c in str(v or '').upper() if c.isalnum())
def ci(d,*names):
    if not isinstance(d,dict): return ''
    x={str(k).lower():v for k,v in d.items()}
    for n in names:
        if n.lower() in x:return x[n.lower()]
    return ''
def name(v):
    if isinstance(v,dict): return str(ci(v,'Name','Value','ProductDescription','Description') or '')
    return str(v or '')
def product(payload): return payload.get('Product') or payload.get('product') or payload

def params(p):
    out={}
    for x in (p.get('Parameters') or p.get('parameters') or []):
        if isinstance(x,dict):
            k=name(ci(x,'ParameterText','Parameter','Name')).strip().lower()
            v=name(ci(x,'ValueText','Value'))
            if k: out[k]=v
    return out

def flatten(v,p=''):
    if isinstance(v,dict):
        for k,x in v.items(): yield from flatten(x,f'{p}.{k}' if p else str(k))
    elif isinstance(v,list):
        for i,x in enumerate(v): yield from flatten(x,f'{p}[{i}]')
    else: yield p,v

def locate_columns(headers):
    mc=pc=None
    for i,h in enumerate(headers,1):
        c=clean(h)
        if c in MFG: mc=i
        if c in MPN: pc=i
    if not mc or not pc: raise ValueError(f'Could not identify Manufacturer and MPN columns: {headers}')
    return mc,pc

def style(ws):
    fill=PatternFill('solid',fgColor='1F4E78')
    for c in ws[1]: c.font=Font(color='FFFFFF',bold=True); c.fill=fill
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        width=min(max(max(len(str(c.value or '')) for c in col)+2,10),55)
        ws.column_dimensions[get_column_letter(col[0].column)].width=width

def run(args):
    src=load_workbook(args.input,data_only=False); ws=src[args.sheet] if args.sheet else src.active
    headers=[c.value for c in ws[1]]; mc,pc=locate_columns(headers)
    rows=[]
    for r in range(max(2,args.start_row),ws.max_row+1):
        mfg=str(ws.cell(r,mc).value or '').strip(); mpn=str(ws.cell(r,pc).value or '').strip()
        if mfg or mpn: rows.append((r,mfg,mpn))
        if args.max_parts and len(rows)>=args.max_parts: break
    settings=Settings.from_env(); print(f'Loaded {len(rows)} parts; DigiKey site={settings.site}, currency={settings.currency}')
    if args.validate_only:return
    client=DigiKeyClient(settings); results=[]; attrs=[]
    manufacturer_catalogue = client.manufacturers(args.force_refresh)
    for n,(row,mfg,mpn) in enumerate(rows,1):
        print(f'[{n}/{len(rows)}] {mfg} {mpn}')
        try:
            resolved = resolve_manufacturer(mfg, manufacturer_catalogue)
            if resolved.manufacturer_id is None:
                raise RuntimeError(
                    f'Manufacturer resolution {resolved.status}: {resolved.reason} '
                    f'(best={resolved.matched_name!r}, confidence={resolved.confidence:.2f})'
                )
            print(
                f'    Manufacturer: {mfg} -> {resolved.matched_name} '
                f'(ID {resolved.manufacturer_id}, confidence {resolved.confidence:.2f})'
            )
            payload=client.details(mpn,resolved.manufacturer_id,args.force_refresh); p=product(payload); pa=params(p)
            mmfg=name(ci(p,'Manufacturer')); mmpn=str(ci(p,'ManufacturerProductNumber','ManufacturerPartNumber','MfrPartNumber') or '')
            mpn_match = norm(mpn)==norm(mmpn)
            returned_mfg_id = ci(ci(p,'Manufacturer'),'Id')
            manufacturer_match = str(returned_mfg_id) == str(resolved.manufacturer_id)
            status='MATCHED' if mpn_match and manufacturer_match else 'REVIEW'
            if status == 'MATCHED':
                reason=(
                    f'Exact normalised MPN with DigiKey manufacturer ID '
                    f'{resolved.manufacturer_id} ({resolved.matched_name}).'
                )
            elif not mpn_match:
                reason='Returned MPN differs; review required.'
            else:
                reason='Returned manufacturer differs from resolved DigiKey manufacturer; review required.'
            results.append([row,mfg,mpn,status,reason,mmfg,mmpn,str(ci(p,'DigiKeyPartNumber','ProductNumber') or ''),name(ci(p,'Description','ProductDescription')),name(ci(p,'DetailedDescription','DetailedProductDescription')),name(ci(p,'Category')),name(ci(p,'Family','ProductFamily')),name(ci(p,'Series')),name(ci(p,'ProductStatus','Status')),name(ci(p,'RoHSStatus','RohsStatus')),pa.get('mounting type',''),pa.get('package / case',pa.get('package/case','')),pa.get('supplier device package',''),pa.get('operating temperature',''),str(ci(p,'DatasheetUrl','DatasheetURL') or ''),str(ci(p,'ProductUrl','ProductURL') or ''),ci(p,'QuantityAvailable'),ci(p,'ManufacturerLeadWeeks'),ci(p,'MinimumOrderQuantity'),datetime.now(timezone.utc).replace(microsecond=0).isoformat()])
            for path,value in flatten(payload): attrs.append([row,mfg,mpn,path,value,'DigiKey Product Information V4'])
        except Exception as e:
            results.append([row,mfg,mpn,'ERROR',str(e),'','','','','','','','','','','','','','','','','','','',datetime.now(timezone.utc).replace(microsecond=0).isoformat()])
    out=Workbook(); e=out.active; e.title='Enriched Parts'
    e.append(['Source Row','Requested Manufacturer','Requested MPN','Match Status','Reason','Matched Manufacturer','Matched MPN','DigiKey Part Number','Description','Detailed Description','Category','Family','Series','Product Status','RoHS Status','Mounting Type','Package / Case','Supplier Device Package','Operating Temperature','Datasheet URL','Product URL','Quantity Available','Lead Weeks','MOQ','Retrieved At UTC'])
    for x in results:e.append(x)
    a=out.create_sheet('All Attributes'); a.append(['Source Row','Requested Manufacturer','Requested MPN','Attribute Path','Attribute Value','Source'])
    for x in attrs:a.append(x)
    r = out.create_sheet('Review Required')
    r.append([cell.value for cell in e[1]])
    for x in results:
        if x[3]!='MATCHED':r.append(x)
    for sh in out.worksheets:style(sh)
    Path(args.output).parent.mkdir(parents=True,exist_ok=True); out.save(args.output); print('Created',args.output)

if __name__=='__main__':
    p=argparse.ArgumentParser(); p.add_argument('--input',required=True); p.add_argument('--output',default='output/AIPN_Enriched.xlsx'); p.add_argument('--sheet'); p.add_argument('--start-row',type=int,default=2); p.add_argument('--max-parts',type=int); p.add_argument('--force-refresh',action='store_true'); p.add_argument('--validate-only',action='store_true'); run(p.parse_args())
